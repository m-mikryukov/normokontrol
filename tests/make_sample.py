#!/usr/bin/env python3
"""Собирает учебный документ с типичными нарушениями оформления.

    python3 tests/make_sample.py tests/samples/bad_coursework.docx
"""

from __future__ import annotations

import struct
import sys
import zlib
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Mm, Pt, RGBColor


def tiny_png(w=120, h=60) -> bytes:
    """PNG без зависимостей: серый прямоугольник."""
    raw = b"".join(b"\x00" + b"\x99\x99\x99" * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def build(path: str):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    s = doc.sections[0]
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(20)

    # титульный лист — не трогаем
    for line in ("Министерство науки и высшего образования", "КУРСОВАЯ РАБОТА",
                 "Анализ процессов", "Москва 2026"):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph("СОДЕРЖАНИЕ").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Введение\t3")
    doc.add_paragraph("1 Теоретическая часть\t4")

    h = doc.add_paragraph("Введение.")          # строчные, точка, не с новой страницы
    h.runs[0].bold = False
    p = doc.add_paragraph("Цель работы — анализ процессов [1]. Схема приведена на рисунке 1, "
                          "данные — в таблице 1, а подробности на рисунке 2 [3].")
    p.paragraph_format.first_line_indent = Cm(1)
    doc.add_paragraph("Задачи работы")          # заголовок без стиля
    doc.paragraphs[-1].runs[0].bold = True

    doc.add_heading("1. Теоретическая часть.", level=1)
    p = doc.add_paragraph("Основной текст раздела с цветным фрагментом. ")
    r = p.add_run("Этот фрагмент красный.")
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    pic = doc.add_paragraph()
    pic.paragraph_format.first_line_indent = Cm(1.25)
    pic.add_run().add_picture(BytesIO(tiny_png()), width=Cm(6))
    doc.add_paragraph("Рис. 1. Схема процесса.")

    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for i, txt in enumerate(("Показатель", "Значение", "Время", "5 мин")):
        t.cell(i // 2, i % 2).text = txt
    doc.add_paragraph("Таблица 1 - Исходные данные")   # под таблицей и с дефисом

    doc.add_paragraph("Заключение").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Работа выполнена.")

    doc.add_paragraph("Список литературы").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Иванов И. И. Процессы. — М.: Наука, 2020. — 100 с.")
    doc.add_paragraph("Петров П. П. Анализ // Вестник. 2021. № 2. С. 1–10.")
    doc.save(path)


def build_dissertation(path: str):
    """Диссертация с типичными нарушениями ГОСТ Р 7.0.11-2011."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(16)                    # больше 14 пт
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(30), Mm(15), Mm(20), Mm(20)

    for line in ("Университет", "На правах рукописи", "Иванов Иван Иванович",
                 "МЕТОДЫ АНАЛИЗА", "Диссертация на соискание учёной степени кандидата наук", "Москва 2026"):
        doc.add_paragraph(line).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph("Оглавление").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Введение\t3")
    doc.add_paragraph("Введение").runs[0].add_break(WD_BREAK.PAGE)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Актуальность темы исследования обусловлена [1], [2].")
    p.paragraph_format.first_line_indent = Cm(2.5)   # больше 17 мм
    doc.add_heading("Глава 1. Обзор методов.", level=1)   # слева, с точкой, не с новой страницы
    doc.add_paragraph("Текст главы со ссылкой на рисунок 1.")
    doc.add_paragraph().add_run().add_picture(BytesIO(tiny_png()), width=Cm(6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Рисунок 1 - Схема")
    doc.add_paragraph("Заключение").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Итоги исследования.")
    doc.add_paragraph("Список литературы").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Петров, П. П. Анализ / П. П. Петров. – Москва : Наука, 2020. – 100 с.")
    doc.add_paragraph("Иванов, И. И. Методы / И. И. Иванов. – Москва : Наука, 2021. – 120 с.")
    doc.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "tests/samples/bad_coursework.docx"
    (build_dissertation if "dissert" in out else build)(out)
