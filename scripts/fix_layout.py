#!/usr/bin/env python3
"""Механические исправления оформления по профилю.

    python3 scripts/fix_layout.py работа.docx -o работа_исправлено.docx [--profile ...]

Все правки вносятся в режиме рецензирования: автор видит их в Word и может
принять или отклонить. Текст автора не меняется; исключение — служебные знаки
оформления (тире и точки в подписях, точка в конце заголовка), они правятся
как видимые вставки/удаления.

Не трогает: титульный лист, содержание, формулы, порядок абзацев, подписи
не на своём месте, ссылки и библиографию — это остаётся в отчёте.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

import nk_common as nk
from audit import ALIGN
from nk_tracked import Revisions, ensure_change_last


class Fixer:
    def __init__(self, doc, profile: dict, track: bool = True, author: str = "Нормоконтроль"):
        self.doc = doc
        self.p = profile
        self.rev = Revisions(doc, author=author, enabled=track)
        self.fonts = nk.Fonts(doc)
        self.blocks = nk.iter_blocks(doc)
        nk.classify(doc, self.blocks, profile)
        self.done: Counter = Counter()
        self.skipped: list[str] = []

    # ------------------------------------------------------------ запуск

    def run(self):
        self.fix_page()
        self.fix_page_numbers()
        for i, b in enumerate(self.blocks):
            if b.kind != "p":
                continue
            role = b.role
            if role in ("body", "list", "bib_entry"):
                self.fix_text_para(b, indent=(role == "body"))
                self.fix_runs(b)
            elif role == "heading_struct":
                self.fix_struct_heading(i, b)
            elif role == "heading_appendix":
                self.fix_appendix_heading(i, b)
            elif role == "heading_section":
                self.fix_section_heading(b)
            elif role in ("caption_figure", "caption_table"):
                self.fix_caption(b, "figure" if role == "caption_figure" else "table")
            elif role == "figure":
                self.fix_figure(b)
            elif role == "table_text":
                self.fix_runs(b, size_min=nk.prof(self.p, "text.table_text_size_pt_min"))
        for part in [self.doc.part] + [r.target_part for r in self.doc.part.rels.values()
                                       if not r.is_external and hasattr(r.target_part, "element")]:
            ensure_change_last(part.element)
        return self

    def note(self, what: str):
        self.done[what] += 1

    # ------------------------------------------------------------ страница

    def fix_page(self):
        margins = nk.prof(self.p, "page.margins_mm", {})
        tol = nk.prof(self.p, "page.tolerance_mm", 0.5)
        for s in self.doc.sections:
            with self.rev.section(s):
                w, h = nk.mm(s.page_width), nk.mm(s.page_height)
                landscape = w > h
                if nk.prof(self.p, "page.size") == "A4":
                    short, long_ = sorted((w, h))
                    if abs(short - 210) > 1 or abs(long_ - 297) > 1:
                        s.page_width, s.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
                        self.note("формат страницы A4")
                if landscape:
                    continue
                for side, need in margins.items():
                    if abs(nk.mm(getattr(s, f"{side}_margin")) - need) > tol:
                        setattr(s, f"{side}_margin", Mm(need))
                        self.note("поля страницы")

    # ------------------------------------------------------------ номера страниц

    def fix_page_numbers(self):
        from audit import Audit

        cfg = nk.prof(self.p, "page_numbers")
        if not cfg:
            return
        probe = Audit(self.doc, self.p)
        found = probe._page_field_paragraphs()
        if not found:
            s = self.doc.sections[0]
            footer = s.footer
            footer.is_linked_to_previous = False
            par = footer.paragraphs[0] if footer.paragraphs and not footer.paragraphs[0].text.strip() \
                else footer.add_paragraph()
            if cfg.get("position", "").endswith("center"):
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            runs = []
            for kind in ("begin", "instr", "separate", "text", "end"):
                r = OxmlElement("w:r")
                if kind in ("begin", "separate", "end"):
                    fc = OxmlElement("w:fldChar")
                    fc.set(qn("w:fldCharType"), kind)
                    r.append(fc)
                elif kind == "instr":
                    it = OxmlElement("w:instrText")
                    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                    it.text = " PAGE "
                    r.append(it)
                else:
                    t = OxmlElement("w:t")
                    t.text = "1"
                    r.append(t)
                par._p.append(r)
                runs.append(r)
            self._font_on_new_runs(runs)
            self.rev.wrap_inserted(runs)
            self.note("номер страницы в нижнем колонтитуле")
        else:
            for in_footer, p_el in found:
                jc = p_el.find(f"{qn('w:pPr')}/{qn('w:jc')}")
                if cfg.get("position", "").endswith("center") and \
                        (jc is None or jc.get(qn("w:val")) != "center") and "\t" not in nk.para_text(p_el):
                    from docx.text.paragraph import Paragraph

                    par = Paragraph(p_el, None)
                    with self.rev.para(par):
                        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.note("номер страницы по центру")
        if cfg.get("title_page_shown") is False and any(b.role == "title" for b in self.blocks):
            s = self.doc.sections[0]
            if not s.different_first_page_header_footer:
                with self.rev.section(s):
                    s.different_first_page_header_footer = True
                self.note("номер на титульном листе скрыт")

    def _font_on_new_runs(self, runs):
        font = nk.prof(self.p, "text.font")
        for r in runs:
            rpr = r.get_or_add_rPr()
            if font:
                rf = rpr.get_or_add_rFonts()
                for attr in ("ascii", "hAnsi", "cs"):
                    rf.set(qn(f"w:{attr}"), font)

    # ------------------------------------------------------------ абзацы

    def _set_align(self, b, align_key: str, label: str):
        need = ALIGN[align_key]
        if (nk.eff_para(self.doc, b.obj, "alignment") or WD_ALIGN_PARAGRAPH.LEFT) != need:
            b.obj.paragraph_format.alignment = need
            self.note(label)

    def _set_indent(self, b, cm: float, label: str):
        ind = nk.eff_para(self.doc, b.obj, "first_line_indent")
        have = 0.0 if ind is None else ind.cm
        if abs(have - cm) > 0.05:
            b.obj.paragraph_format.first_line_indent = Cm(cm)
            self.note(label)

    def fix_text_para(self, b, indent=True):
        t = self.p.get("text", {})
        with self.rev.para(b.obj):
            if "line_spacing" in t:
                ls = nk.eff_para(self.doc, b.obj, "line_spacing")
                if not isinstance(ls, float) or abs(ls - t["line_spacing"]) > 0.01:
                    b.obj.paragraph_format.line_spacing = t["line_spacing"]
                    self.note("межстрочный интервал")
            if "alignment" in t:
                self._set_align(b, t["alignment"], "выравнивание текста")
            if indent and "first_line_indent_cm" in t:
                self._set_indent(b, t["first_line_indent_cm"], "абзацный отступ")

    def fix_runs(self, b, size_min="profile", bold=None, caps=False):
        t = self.p.get("text", {})
        font = t.get("font")
        if size_min == "profile":
            size_min = t.get("size_pt_min")
        for r in nk.text_runs(b):
            with self.rev.run(r):
                rpr = r.get_or_add_rPr()
                if font and self.fonts.font(r, b.obj) != font:
                    rf = rpr.get_or_add_rFonts()
                    for attr in ("asciiTheme", "hAnsiTheme", "cstheme"):
                        rf.attrib.pop(qn(f"w:{attr}"), None)
                    for attr in ("ascii", "hAnsi", "cs"):
                        rf.set(qn(f"w:{attr}"), font)
                    self.note("шрифт")
                vert = rpr.find(qn("w:vertAlign"))
                if size_min and vert is None and self.fonts.size_pt(r, b.obj) < size_min:
                    rpr.get_or_add_sz().val = Pt(size_min)
                    szcs = rpr.find(qn("w:szCs"))
                    if szcs is None:
                        szcs = OxmlElement("w:szCs")
                        rpr.find(qn("w:sz")).addnext(szcs)
                    szcs.set(qn("w:val"), str(int(size_min * 2)))
                    self.note("кегль")
                if t.get("color") == "black" and not nk.is_black(self.fonts.color(r, b.obj)):
                    c = rpr.get_or_add_color()
                    c.val = RGBColor(0, 0, 0)
                    for attr in ("themeColor", "themeTint", "themeShade"):
                        c.attrib.pop(qn(f"w:{attr}"), None)
                    self.note("цвет текста")
                if bold and not self.fonts.flag(r, b.obj, "b"):
                    rpr.get_or_add_b().val = True
                    self.note("полужирный заголовок")
                if caps and not self.fonts.flag(r, b.obj, "caps") and run_has_lower(r):
                    rpr.get_or_add_caps().val = True
                    self.note("прописные буквы в заголовке")

    # ------------------------------------------------------------ заголовки

    def _remove_trailing_dot(self, b, label):
        text = b.text
        stripped = text.rstrip()
        if stripped.endswith(".") and not stripped.endswith(".."):
            pos = len(stripped) - 1
            if self.rev.replace_span(b.el, pos, pos + 1, ""):
                self.note(label)
            else:
                self.skipped.append(f"{b.where()}: не удалось убрать точку автоматически")

    def _ensure_new_page(self, i, b, label):
        if not nk.starts_new_page(self.doc, self.blocks, i):
            b.obj.paragraph_format.page_break_before = True
            self.note(label)

    def fix_struct_heading(self, i, b):
        cfg = nk.prof(self.p, "headings.structural_elements", {})
        with self.rev.para(b.obj):
            if "alignment" in cfg:
                self._set_align(b, cfg["alignment"], "выравнивание заголовка")
            self._set_indent(b, 0, "отступ заголовка")
            if cfg.get("new_page"):
                self._ensure_new_page(i, b, "заголовок с новой страницы")
        self.fix_runs(b, bold=cfg.get("bold"), caps=cfg.get("case") == "upper")
        if cfg.get("trailing_dot") is False:
            self._remove_trailing_dot(b, "точка в конце заголовка")

    def fix_appendix_heading(self, i, b):
        cfg = nk.prof(self.p, "appendices", {})
        with self.rev.para(b.obj):
            if cfg.get("alignment"):
                self._set_align(b, cfg["alignment"], "выравнивание заголовка")
                self._set_indent(b, 0, "отступ заголовка")
            if cfg.get("new_page"):
                self._ensure_new_page(i, b, "приложение с новой страницы")
        self.fix_runs(b)

    def fix_section_heading(self, b):
        cfg = nk.prof(self.p, "headings.sections", {})
        with self.rev.para(b.obj):
            if "alignment" in cfg:
                self._set_align(b, cfg["alignment"], "выравнивание заголовка")
            if "first_line_indent_cm" in cfg:
                self._set_indent(b, cfg["first_line_indent_cm"], "отступ заголовка")
        self.fix_runs(b, bold=cfg.get("bold"))
        if cfg.get("trailing_dot") is False:
            self._remove_trailing_dot(b, "точка в конце заголовка")
        if cfg.get("number_trailing_dot") is False:
            m = re.match(r"^(\s*\d+(?:\.\d+)*)\.\s", b.text)
            if m:
                if self.rev.replace_span(b.el, m.end(1), m.end(1) + 1, ""):
                    self.note("точка после номера раздела")
                else:
                    self.skipped.append(f"{b.where()}: точку после номера убрать вручную")

    # ------------------------------------------------------------ подписи и рисунки

    def fix_caption(self, b, kind):
        cfg = nk.prof(self.p, "figures" if kind == "figure" else "tables", {})
        with self.rev.para(b.obj):
            if "caption_alignment" in cfg:
                self._set_align(b, cfg["caption_alignment"], "выравнивание подписи")
            if "caption_first_line_indent_cm" in cfg:
                self._set_indent(b, cfg["caption_first_line_indent_cm"], "отступ подписи")
        self.fix_runs(b)

        # текст подписи: правим справа налево, чтобы смещения не поехали
        lead = len(b.text) - len(b.text.lstrip())
        m = nk.match_caption(b.text, kind)
        title = m.group("title")
        edits = []
        if cfg.get("caption_trailing_dot") is False:
            tr = title.rstrip()
            if tr.endswith(".") and not tr.endswith("..") and not re.search(r"\b\w{1,2}\.$", tr):
                pos = lead + m.start("title") + len(tr) - 1
                edits.append((pos, pos + 1, "", "точка в конце подписи"))
        if m.group("sep") != " — ":
            edits.append((lead + m.start("sep"), lead + m.end("sep"), " — ", "тире в подписи"))
        if m.group("sp") != " ":
            edits.append((lead + m.start("sp"), lead + m.end("sp"), " ", "пробел в подписи"))
        if m.group("label") != nk.CANON_LABEL[kind]:
            edits.append((lead + m.start("label"), lead + m.end("label"), nk.CANON_LABEL[kind],
                          "слово «Рисунок»/«Таблица» в подписи"))
        for start, end, new, label in sorted(edits, reverse=True):
            if self.rev.replace_span(b.el, start, end, new):
                self.note(label)
            else:
                self.skipped.append(f"{b.where()}: {label} — поправить вручную (поле или сложная разметка)")

    def fix_figure(self, b):
        cfg = nk.prof(self.p, "figures", {})
        if not cfg.get("image_alignment"):
            return
        with self.rev.para(b.obj):
            self._set_align(b, cfg["image_alignment"], "выравнивание рисунка")
            ind = nk.eff_para(self.doc, b.obj, "first_line_indent")
            if ind is not None and abs(ind.cm) > 0.05:
                b.obj.paragraph_format.first_line_indent = Cm(0)
                self.note("отступ у рисунка")

    # ------------------------------------------------------------ итог

    def summary(self) -> str:
        lines = [f"Внесено правок: {self.rev.count}"
                 + (" (в режиме рецензирования)" if self.rev.enabled else " (без рецензирования)") + "."]
        for what, n in self.done.most_common():
            lines.append(f"- {what}: {n}")
        if self.skipped:
            lines.append("")
            lines.append("Не удалось исправить автоматически:")
            lines += [f"- {s}" for s in self.skipped]
        return "\n".join(lines) + "\n"


def run_has_lower(r_el) -> bool:
    return any(ch.islower() for ch in nk.run_text(r_el))


def main(argv=None):
    ap = argparse.ArgumentParser(description="Механические исправления оформления .docx")
    ap.add_argument("docx")
    ap.add_argument("-o", "--output", required=True)
    ap.add_argument("--profile", default="gost-7.32-2017", help="id профиля или путь к .yaml")
    ap.add_argument("--no-track", action="store_true", help="править без режима рецензирования")
    ap.add_argument("--author", default="Нормоконтроль", help="автор правок в рецензировании")
    args = ap.parse_args(argv)

    if args.output == args.docx:
        ap.error("исходный файл не перезаписываем: укажите другой -o")
    profile = nk.load_profile(args.profile)
    doc = nk.open_docx(args.docx)
    fixer = Fixer(doc, profile, track=not args.no_track, author=args.author).run()
    doc.save(args.output)
    sys.stdout.write(fixer.summary())
    return 0


if __name__ == "__main__":
    sys.exit(main())
